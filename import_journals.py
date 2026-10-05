"""Import an explicitly supplied JCR journal list; never publish workbook metrics."""
from __future__ import annotations
import argparse, collections, datetime as dt, hashlib, json, re, urllib.parse, urllib.request
from pathlib import Path

ROOT=Path(__file__).resolve().parent
INDEXES={'SSCI','SCIE'}
def normalized(value):return re.sub(r'[^a-z0-9]','',str(value).casefold().replace('&','and'))
def issns(values):return sorted({str(v).strip().upper() for v in values if re.fullmatch(r'\d{4}-\d{3}[\dXx]',str(v).strip())})
def publisher_group(name):
    n=name.upper()
    for group,terms in [('Nature Portfolio',['NATURE PORTFOLIO']),('Taylor & Francis',['TAYLOR','ROUTLEDGE','DOVE MEDICAL']),('Elsevier',['ELSEVIER','PERGAMON','CHURCHILL LIVINGSTONE']),('Springer Nature',['SPRINGER','BMC','PALGRAVE']),('Wiley',['WILEY']),('Sage',['SAGE PUBLICATIONS']),('IEEE',['IEEE']),('IFETS',['INT FORUM EDUCATIONAL']),('APA',['AMER PSYCHOLOGICAL','AMERICAN PSYCHOLOGICAL'])]:
        if any(t in n for t in terms):return group
    return name

def workbook_rows(path):
    from openpyxl import load_workbook
    workbook=load_workbook(path,read_only=True,data_only=True);rows=[];counts={}
    for sheet in workbook:
        iterator=sheet.iter_rows(values_only=True);headers=next(iterator);counts[sheet.title]=0
        for rownum,values in enumerate(iterator,2):
            row=dict(zip(headers,values));edition=str(row.get('Edition') or '').strip().upper()
            if edition not in INDEXES:continue
            name=str(row.get('Journal name') or '').strip();identifiers=issns([row.get('ISSN'),row.get('eISSN')])
            if not name or not identifiers:raise ValueError(f'Missing journal identity at {sheet.title}:{rownum}')
            rows.append({'name':name,'issns':identifiers,'publisher':str(row.get('Publisher') or ''),'edition':edition,'category':row.get('Category'),'sheet':sheet.title,'row':rownum});counts[sheet.title]+=1
    return rows,counts

def merge_rows(existing,rows,filename,asof,period='2026-06'):
    journals=json.loads(json.dumps(existing));added=0;matched=set()
    for row in rows:
        matches=[j for j in journals if set(j.get('issns',[]))&set(row['issns']) or (not j.get('issns') and normalized(j['name'])==normalized(row['name']))]
        if len(matches)>1:raise ValueError('Ambiguous existing journal: '+row['name'])
        if matches:j=matches[0]
        else:
            j={'id':hashlib.sha256(row['issns'][0].encode()).hexdigest()[:12],'name':row['name'],'publisher':publisher_group(row['publisher']),'issns':[],'indexing':[],'verified':True,'checked_at':asof,'evidence_url':None,'evidence_type':'user_jcr','journal_url':None}
            journals.append(j);added+=1
        matched.add(j['id']);j['issns']=sorted(set(j.get('issns',[]))|set(row['issns']))
        j.setdefault('aliases',[])
        if row['name']!=j['name'] and row['name'] not in j['aliases']:j['aliases'].append(row['name'])
        if j.get('registry_provenance',{}).get('file')!=filename or j.get('registry_provenance',{}).get('period')!=period:
            j['registry_provenance']={'kind':'user_jcr','file':filename,'period':period,'imported_at':asof,'rows':[]}
        entry=j['registry_provenance']
        ref={'sheet':row['sheet'],'row':row['row'],'edition':row['edition'],'category':row['category']}
        if ref not in entry['rows']:entry['rows'].append(ref)
        j['listed_indexing']=sorted({r['edition'] for r in entry['rows']})
        # Preserve live publisher evidence for existing journals. A spreadsheet
        # import must never overwrite a later official downgrade or delisting.
        if j.get('evidence_type')=='user_jcr':j['indexing']=j['listed_indexing']
        j.setdefault('scope_topics',[])
        topic='心理学' if row['sheet']=='Psychology' else '教育学'
        if topic not in j['scope_topics']:j['scope_topics'].append(topic)
    return journals,{'added':added,'unique_imported':len(matched),'existing_retained':len(existing),'registry_total':len(journals)}

def directory_lookup(identifiers):
    params=urllib.parse.urlencode({'filter':'issn:'+'|'.join(identifiers),'per_page':100,'select':'id,display_name,issn,homepage_url,host_organization_name,type'})
    req=urllib.request.Request('https://api.openalex.org/sources?'+params,headers={'User-Agent':'CallAtlas/1.0 journal-directory resolver'})
    with urllib.request.urlopen(req,timeout=30) as response:return json.load(response)['results']

def homepage(raw):
    if not raw:return None
    p=urllib.parse.urlsplit(raw)
    if p.scheme not in ('http','https') or not p.hostname or p.username or p.password:return None
    host=p.hostname.lower();path=p.path
    if host in ('onlinelibrary.wiley.com','www3.interscience.wiley.com'):
        match=re.search(r'(\d{4})-?(\d{3}[\dXx])',path)
        if match:return 'https://onlinelibrary.wiley.com/journal/'+''.join(match.groups()).lower()
    if host in ('www.tandfonline.com','tandfonline.com'):
        m=re.search(r'/(?:toc|loi|journals)/([a-z0-9]+)',path,re.I)
        code=(urllib.parse.parse_qs(p.query).get('journalCode') or [None])[0]
        if m or code:return 'https://www.tandfonline.com/journals/'+(m.group(1) if m else code).lower()
    if host in ('journals.sagepub.com','www.sagepub.com') and path.startswith('/home/'):return 'https://journals.sagepub.com'+path.rstrip('/')
    if re.fullmatch(r'[a-z]{2,5}\.sagepub\.com',host) and host.split('.')[0] not in ('www','us','uk'):
        return 'https://journals.sagepub.com/home/'+host.split('.')[0]
    if host in ('link.springer.com','www.springer.com'):
        m=re.search(r'/(?:journal/)?(\d+)$',path)
        if m:return 'https://link.springer.com/journal/'+m.group(1)
    if host in ('journals.elsevier.com','www.journals.elsevier.com'):
        return 'https://www.sciencedirect.com/journal/'+path.strip('/').split('/')[0]
    return urllib.parse.urlunsplit(('https',p.netloc,path,p.query,''))

def resolve(journals,directory):
    outcomes=collections.Counter()
    for j in journals:
        if j.get('journal_url'):outcomes['retained']+=1;continue
        matches={d['id']:d for d in directory if d.get('type')=='journal' and set(d.get('issn') or [])&set(j['issns'])}
        if len(matches)!=1:j['directory_status']='not_found' if not matches else 'ambiguous';outcomes[j['directory_status']]+=1;continue
        d=next(iter(matches.values()));url=homepage(d.get('homepage_url'))
        j['directory_provenance']={'provider':'OpenAlex','id':d['id'],'issn_match':True,'raw_homepage':d.get('homepage_url')}
        if not url:j['directory_status']='no_homepage';outcomes['no_homepage']+=1;continue
        j['journal_url']=url;j['directory_status']='issn_matched';outcomes['resolved']+=1
        if normalized(d['display_name'])!=normalized(j['name']):j.setdefault('aliases',[]).append(d['display_name'])
        j['evidence_candidate_url']=url
        if '/journals/' in url and 'tandfonline.com' in url:j['evidence_candidate_url']=url+'/about-this-journal'
        if '/journal/' in url and 'onlinelibrary.wiley.com' in url:j['evidence_candidate_url']=url.replace('/journal/','/page/journal/')+'/homepage/productinformation.html'
        if '/journal/' in url and 'sciencedirect.com' in url:j['evidence_candidate_url']=url+'/about/insights'
    return dict(outcomes)

def main():
    parser=argparse.ArgumentParser();parser.add_argument('workbook');parser.add_argument('--directory-cache',required=True);parser.add_argument('--period',help='JCR snapshot month YYYY-MM; otherwise inferred from _MM_YYYY filename');args=parser.parse_args()
    registry=json.loads((ROOT/'journals.json').read_text(encoding='utf-8'));asof=dt.date.today().isoformat()
    month=re.search(r'_(\d{2})_(\d{4})',Path(args.workbook).name);period=args.period or (month[2]+'-'+month[1] if month else None)
    if not period:parser.error('Specify the actual JCR snapshot month with --period YYYY-MM')
    dt.date.fromisoformat(period+'-01')
    rows,counts=workbook_rows(args.workbook);journals,summary=merge_rows(registry['journals'],rows,Path(args.workbook).name,asof,period)
    directory=json.loads(Path(args.directory_cache).read_text(encoding='utf-8'));resolved=resolve(journals,directory)
    registry.update(journals=journals,import_summary={**summary,'sheet_rows':counts,'imported_at':asof,'source_file':Path(args.workbook).name,'source_period':period,'directory':resolved})
    (ROOT/'journals.json').write_text(json.dumps(registry,ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n')
    print(json.dumps(registry['import_summary'],ensure_ascii=False))
if __name__=='__main__':main()
